import os
import re

db_js_path = os.path.join(os.path.dirname(__file__), '..', 'config', 'db.js')
with open(db_js_path, 'r', encoding='utf-8') as f:
    db_js = f.read()

# Just reuse the generated final_sql from the previous script run
# Wait, I need to generate it again. I'll just copy the code over and fix the target.
import openpyxl
from datetime import datetime

file_path = r'C:\Users\asada\OneDrive\Desktop\AI Mentorship Program Attendance & Daily Task.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)

projects_dict = {}
users_dict = {}

def clean_id(name):
    return re.sub(r'[^a-zA-Z0-9]', '', name).lower()

for sheet in wb.sheetnames:
    ws = wb[sheet]
    headers = [cell.value for cell in ws[1]]
    if not headers or 'Project Name' not in headers:
        for r in range(1, 5):
            headers = [cell.value for cell in ws[r]]
            if 'Project Name' in headers:
                start_row = r + 1
                break
        else:
            continue
    else:
        start_row = 2

    if 'Project Name' not in headers:
        continue
        
    proj_idx = headers.index('Project Name')
    student_idx = headers.index('Student Name')
    work_idx = headers.index('Work Completed Today') if 'Work Completed Today' in headers else -1
    date_idx = headers.index('Date') if 'Date' in headers else -1

    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if not row[proj_idx] or not row[student_idx]:
            continue
        proj = str(row[proj_idx]).strip()
        student = str(row[student_idx]).strip()
        work = str(row[work_idx]).strip() if work_idx != -1 and row[work_idx] else ''
        date_raw = row[date_idx] if date_idx != -1 else ''
        
        date_str = ''
        if isinstance(date_raw, datetime):
            date_str = date_raw.strftime('%Y-%m-%d')
        elif date_raw:
            date_raw = str(date_raw).strip()
            try:
                date_str = datetime.strptime(date_raw, '%d/%m/%Y').strftime('%Y-%m-%d')
            except:
                try:
                    date_str = datetime.strptime(date_raw, '%Y-%m-%d %H:%M:%S').strftime('%Y-%m-%d')
                except:
                    date_str = '2026-08-24'
        else:
            date_str = '2026-08-24'
            
        if proj.lower() == 'project name' or proj == 'None' or not proj:
            continue
            
        if student not in users_dict:
            users_dict[student] = f"usr-mentor-{clean_id(student)}"
            
        if proj not in projects_dict:
            projects_dict[proj] = {'id': f"proj-mentor-{clean_id(proj)}", 'students': set(), 'tasks': []}
            
        projects_dict[proj]['students'].add(student)
        
        if work and work.lower() != 'none':
            projects_dict[proj]['tasks'].append({
                'assignee': student,
                'task': work,
                'date': date_str
            })

users_sql = []
for student, uid in users_dict.items():
    users_sql.append(f"          ('{uid}', '{student.replace(chr(39), '')}', '{uid}@qarc.ai', 'Intern', '/default-avatar.png')")
users_block = "INSERT INTO users (id, name, email, role, \\\"avatarUrl\\\") VALUES \\n" + ",\\n".join(users_sql) + "\\nON CONFLICT (id) DO NOTHING;" if users_sql else ""

programs_block = "INSERT INTO programs (id, name, description, start_date, end_date) VALUES ('prog-mentor-b1', 'Mentorship Program Batch 1', 'AI Mentorship Program', '2026-08-24', '2026-09-20') ON CONFLICT (id) DO NOTHING;"

projects_sql = []
pp_sql = []
pm_sql = []
tasks_sql = []

task_counter = 1
pm_counter = 1

for proj, pdata in projects_dict.items():
    pid = pdata['id']
    safe_proj = proj.replace(chr(39), "''")
    projects_sql.append(f"          ('{pid}', '{safe_proj}', 'AI Mentorship', 'Mentorship Project', '{safe_proj}', 'In Progress', 'High', '2026-08-24', '2026-09-20', 50, 'prog-mentor-b1', 'usr-pm-asad', 'usr-pm-asad', 'usr-pm-asad')")
    pp_sql.append(f"          ('pp-{pid}', 'prog-mentor-b1', '{safe_proj}', 'AI Mentorship', 'Mentorship Project', '{safe_proj}', 'In Progress', 'High', '2026-08-24', '2026-09-20', 'usr-pm-asad', 'usr-pm-asad')")
    
    for student in pdata['students']:
        uid = users_dict[student]
        pm_sql.append(f"          ('pm-m-{pm_counter}', '{pid}', '{uid}', 'Intern')")
        pm_counter += 1
        
    for t in pdata['tasks']:
        uid = users_dict[t['assignee']]
        safe_task = t['task'].replace(chr(39), "''").replace("\\n", " ")
        if len(safe_task) > 200:
            safe_task = safe_task[:197] + "..."
        date = t['date']
        tasks_sql.append(f"          ('tsk-m-{task_counter}', '{safe_task}', '{safe_task}', 'Done', 'Medium', '{date}', '{date}', '{pid}', '{uid}', '{uid}', 'usr-pm-asad', '{date} 23:59:59')")
        task_counter += 1

projects_block = "INSERT INTO projects (id, name, domain, about_title, about_description, status, priority, start_date, deadline, progress, program_id, pm_id, tl_id, created_by) VALUES \\n" + ",\\n".join(projects_sql) + "\\nON CONFLICT (id) DO NOTHING;" if projects_sql else ""
pp_block = "INSERT INTO program_projects (id, program_id, project_name, domain, about_title, about_description, status, priority, start_date, deadline, pm_id, tl_id) VALUES \\n" + ",\\n".join(pp_sql) + "\\nON CONFLICT (id) DO NOTHING;" if pp_sql else ""
pm_block = "INSERT INTO project_members (id, project_id, user_id, role) VALUES \\n" + ",\\n".join(pm_sql) + "\\nON CONFLICT (id) DO NOTHING;" if pm_sql else ""
tasks_block = "INSERT INTO tasks (id, name, description, status, priority, start_date, due_date, project_id, assignee_id, assigned_to, created_by, completed_at) VALUES \\n" + ",\\n".join(tasks_sql) + "\\nON CONFLICT (id) DO NOTHING;" if tasks_sql else ""

final_sql = f'''
        // MENTORSHIP BATCH 1 INJECTION
        await client.query(
          {users_block}
          {programs_block}
          {projects_block}
          {pp_block}
          {pm_block}
          {tasks_block}
        );
'''

target_regex = re.compile(r"console\.log\(\"[^\"]*QARC Database schema")
match = target_regex.search(db_js)
if match:
    db_js = db_js[:match.start()] + final_sql + "      " + db_js[match.start():]
    with open(db_js_path, 'w', encoding='utf-8') as f:
        f.write(db_js)
    print(f"Successfully injected Mentorship Program Batch 1 with {len(projects_dict)} projects and {task_counter-1} tasks!")
else:
    print("Could not find the injection point in db.js")
